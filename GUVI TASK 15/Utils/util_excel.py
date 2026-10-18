from openpyxl import load_workbook
from datetime import datetime

def get_test_data(file_path, sheet_name):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]
    data = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        data.append((row[1], row[2]))
    return data

def write_result(file_path, sheet_name, row, tester, actual_result):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    now = datetime.now()
    sheet.cell(row=row, column=4).value = now.strftime("%Y-%m-%d")   # Date
    sheet.cell(row=row, column=5).value = now.strftime("%H:%M:%S")   # Time
    sheet.cell(row=row, column=6).value = tester
    sheet.cell(row=row, column=7).value = actual_result

    workbook.save(file_path)

